# -*- coding: utf-8 -*-
"""
Vuelve a crear en el lector los usuarios de un Excel de revision.

Es el reverso de borrar_usuarios_lector.py: toma los que tengan SI en la
columna "Borrar" y los da de alta de nuevo en el equipo, con su ID y su nombre.

OJO CON LA BIOMETRIA: esto restaura el usuario, NO su huella ni su rostro. Al
borrarlo del lector se fue tambien su biometria, y salvo que estuviera
respaldada en el panel no hay de donde sacarla. Cada persona tiene que volver a
enrolar el dedo. Si la huella si esta respaldada en el panel, el propio panel la
vuelve a subir cuando sincroniza.

La vigencia se pone desde hoy hasta hoy + los anios configurados en el panel,
porque la original no quedo guardada en ningun lado.

Simula por defecto.

    python instalacion\\restaurar_usuarios_lector.py --excel usuarios_a_borrar_254.xlsx --ip 192.168.88.254
    python instalacion\\restaurar_usuarios_lector.py --excel usuarios_a_borrar_254.xlsx --ip 192.168.88.254 --aplicar
"""

import argparse
import logging
import os
import sys
from ctypes import sizeof

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, BASE)
os.chdir(BASE)
try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

logging.basicConfig(level=logging.WARNING, format="%(message)s", stream=sys.stderr)
import panel_personas as P  # noqa: E402
logging.getLogger().handlers = [logging.StreamHandler(sys.stderr)]
logging.getLogger().setLevel(logging.WARNING)

from openpyxl import load_workbook  # noqa: E402
from SDK_Struct import (NET_IN_LOGIN_WITH_HIGHLEVEL_SECURITY,  # noqa: E402
                        NET_OUT_LOGIN_WITH_HIGHLEVEL_SECURITY)
from SDK_Enum import EM_LOGIN_SPAC_CAP_TYPE  # noqa: E402


def leer_excel(ruta):
    wb = load_workbook(ruta, read_only=True)
    ws = wb[wb.sheetnames[0]]
    filas = list(ws.values)
    cab = [str(c or "").strip().lower() for c in filas[0]]
    try:
        i_id = cab.index("id en el lector")
        i_nom = cab.index("nombre en el lector")
        i_bor = cab.index("borrar")
    except ValueError:
        raise SystemExit("El Excel no tiene las columnas esperadas: %s" % cab)
    elegidos = []
    for f in filas[1:]:
        if not f or f[i_id] is None:
            continue
        if str(f[i_bor] or "").strip().upper() != "SI":
            continue
        elegidos.append((str(f[i_id]).strip(), str(f[i_nom] or "").strip()))
    return elegidos


def entrar(ip):
    dev = next((d for d in P.DEVICES if d["ip"] == ip), None)
    if not dev:
        raise SystemExit("El lector %s no esta en config.json" % ip)
    ent = NET_IN_LOGIN_WITH_HIGHLEVEL_SECURITY()
    ent.dwSize = sizeof(ent)
    ent.szIP = ip.encode()
    ent.nPort = int(dev.get("port", 37777))
    ent.szUserName = str(dev["user"]).encode()
    ent.szPassword = str(dev["password"]).encode()
    ent.emSpecCap = EM_LOGIN_SPAC_CAP_TYPE.TCP
    sal = NET_OUT_LOGIN_WITH_HIGHLEVEL_SECURITY()
    sal.dwSize = sizeof(sal)
    P.client.InitEx(None)
    login_id, _i, err = P.client.LoginWithHighLevelSecurity(ent, sal)
    if not login_id:
        raise SystemExit("No pude entrar a %s: %s" % (ip, err))
    return int(login_id)


def main():
    ap = argparse.ArgumentParser(description="Restaura usuarios en un lector desde un Excel")
    ap.add_argument("--excel", required=True)
    ap.add_argument("--ip", required=True)
    ap.add_argument("--aplicar", action="store_true", help="sin esto solo simula")
    args = ap.parse_args()

    gente = leer_excel(args.excel)
    desde, hasta = P.vigencia_por_defecto()
    d_txt, h_txt = desde.strftime("%Y-%m-%d"), hasta.strftime("%Y-%m-%d")

    print("=" * 74)
    print("  %s   lector %s" % ("RESTAURANDO" if args.aplicar else "SIMULACION - no toca nada",
                                args.ip))
    print("=" * 74)
    print("  usuarios a restaurar : %d" % len(gente))
    print("  vigencia que se pone : %s  a  %s" % (d_txt, h_txt))
    print("  biometria            : NO se restaura, hay que enrolar el dedo de nuevo")

    if not gente:
        print("\n  No hay nadie marcado con SI.")
        return 0

    if not args.aplicar:
        print("\n  primeros 10:")
        for uid, nom in gente[:10]:
            print("     %-12s %s" % (uid, nom[:40]))
        print("\n" + "-" * 74)
        print("  SIMULACION. No se creo nada. Para hacerlo, agregar  --aplicar")
        return 0

    login_id = entrar(args.ip)
    creados, ya_estaban, fallaron = [], [], []
    try:
        for n, (uid, nom) in enumerate(gente, start=1):
            if P.sdk_consultar(login_id, uid) is not None:
                ya_estaban.append(uid)
                continue
            # sdk_alta espera datetime, no texto: adentro hace desde.year
            ok, msg = P.sdk_alta(login_id, uid, nom, desde, hasta)
            (creados if ok else fallaron).append((uid, nom, msg if not ok else ""))
            if n % 25 == 0:
                print("   %d/%d  (creados %d, fallaron %d)"
                      % (n, len(gente), len(creados), len(fallaron)))
    finally:
        try:
            P.client.Logout(login_id)
        except Exception:
            pass

    print("\n" + "-" * 74)
    print("  creados       : %d" % len(creados))
    print("  ya existian   : %d" % len(ya_estaban))
    print("  fallaron      : %d" % len(fallaron))
    for uid, nom, msg in fallaron[:15]:
        print("     %-12s %-28s %s" % (uid, nom[:28], str(msg)[:40]))

    print("\n  verificando contra el lector...")
    login_id = entrar(args.ip)
    faltan = []
    try:
        for uid, nom, _m in creados:
            if P.sdk_consultar(login_id, uid) is None:
                faltan.append(uid)
    finally:
        try:
            P.client.Logout(login_id)
        except Exception:
            pass
    if faltan:
        print("  ATENCION: %d no aparecen en el equipo pese al alta: %s"
              % (len(faltan), ", ".join(faltan[:15])))
    else:
        print("  verificado: los %d creados estan en el equipo." % len(creados))
    print()
    print("  RECORDAR: estan sin huella ni rostro. No pueden fichar hasta enrolarlos.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
